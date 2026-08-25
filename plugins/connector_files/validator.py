"""Deterministic, bounded CSV/XLSX validation against a published template."""

from __future__ import annotations

from collections.abc import Iterable
import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
from io import BytesIO, StringIO
import json
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from packages.contracts.source_intake import SourceIntakeFailure
from packages.data_documentation.domain.model import (
    FileImportTemplateVersion,
    RejectedRow,
    TemplateColumn,
    ValidatedFileImport,
)


def _schema_fingerprint(template: FileImportTemplateVersion) -> str:
    payload = [
        {
            "header": column.source_header,
            "role": column.field_role,
            "type": column.data_type,
            "required": column.required,
        }
        for column in template.columns
    ]
    return hashlib.sha256(
        json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def _headers(template: FileImportTemplateVersion, actual: Iterable[object]) -> None:
    normalized = tuple("" if item is None else str(item).strip() for item in actual)
    expected = tuple(column.source_header for column in template.columns)
    if normalized != expected or len(normalized) != len(set(normalized)):
        raise SourceIntakeFailure("FILE_STRUCTURE_REJECTED")


def _formula_like(value: str, column: TemplateColumn) -> bool:
    if not value:
        return False
    if value[0] in {"=", "+", "@"}:
        return True
    return value[0] == "-" and column.data_type == "string"


def _parse(value: object, column: TemplateColumn, template: FileImportTemplateVersion) -> object:
    if value is None or (isinstance(value, str) and not value.strip()):
        if column.required:
            raise ValueError("required")
        return None
    if isinstance(value, str):
        value = value.strip()
        if _formula_like(value, column):
            raise SourceIntakeFailure("FILE_FORMULA_REJECTED")
    if column.data_type == "string":
        return str(value)
    if column.data_type == "integer":
        return int(str(value))
    if column.data_type == "decimal":
        raw = str(value)
        if template.decimal_separator == ",":
            raw = raw.replace(".", "").replace(",", ".")
        return Decimal(raw)
    if column.data_type == "boolean":
        if isinstance(value, bool):
            return value
        normalized = str(value).casefold()
        if normalized in {"true", "1", "yes"}:
            return True
        if normalized in {"false", "0", "no"}:
            return False
        raise ValueError("boolean")
    if column.data_type == "date":
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return datetime.strptime(str(value), template.date_format).date()
    if column.data_type == "datetime":
        if isinstance(value, datetime):
            return value
        return datetime.strptime(str(value), template.date_format)
    raise ValueError("type")


def _typed_rows(
    template: FileImportTemplateVersion,
    rows: Iterable[tuple[int, tuple[object, ...]]],
) -> tuple[tuple[dict[str, object], ...], tuple[RejectedRow, ...]]:
    accepted: list[dict[str, object]] = []
    rejected: list[RejectedRow] = []
    for row_number, values in rows:
        if len(values) != len(template.columns):
            code = "COLUMN_COUNT_MISMATCH"
        else:
            try:
                accepted.append(
                    {
                        column.field_role: _parse(value, column, template)
                        for column, value in zip(template.columns, values, strict=True)
                    }
                )
                continue
            except SourceIntakeFailure:
                raise
            except (ValueError, TypeError, InvalidOperation, OverflowError):
                code = "TYPE_OR_REQUIRED_VALUE_INVALID"
        if template.error_policy == "reject_file":
            raise SourceIntakeFailure("FILE_VALIDATION_FAILED")
        rejected.append(RejectedRow(row_number=row_number, code=code))
    if len(accepted) + len(rejected) > template.row_limit:
        raise SourceIntakeFailure("FILE_ROW_LIMIT_EXCEEDED")
    return tuple(accepted), tuple(rejected)


def _csv_rows(
    template: FileImportTemplateVersion, content: bytes
) -> tuple[tuple[dict[str, object], ...], tuple[RejectedRow, ...]]:
    if content.startswith(b"PK"):
        raise SourceIntakeFailure("FILE_MEDIA_TYPE_MISMATCH")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise SourceIntakeFailure("FILE_ENCODING_REJECTED") from exc
    reader = csv.reader(StringIO(text, newline=""))
    try:
        headers = next(reader)
    except StopIteration as exc:
        raise SourceIntakeFailure("FILE_STRUCTURE_REJECTED") from exc
    _headers(template, headers)
    rows = ((number, tuple(values)) for number, values in enumerate(reader, start=2))
    return _typed_rows(template, rows)


def _xlsx_rows(
    template: FileImportTemplateVersion, content: bytes
) -> tuple[tuple[dict[str, object], ...], tuple[RejectedRow, ...]]:
    try:
        with ZipFile(BytesIO(content)) as archive:
            names = set(archive.namelist())
            expanded = sum(item.file_size for item in archive.infolist())
            if expanded > template.file_size_limit * 20:
                raise SourceIntakeFailure("FILE_EXPANDED_SIZE_REJECTED")
            forbidden = (
                "vbaProject.bin",
                "xl/externalLinks/",
                "xl/embeddings/",
            )
            if any(any(marker in name for marker in forbidden) for name in names):
                raise SourceIntakeFailure("FILE_ACTIVE_CONTENT_REJECTED")
    except BadZipFile as exc:
        raise SourceIntakeFailure("FILE_MEDIA_TYPE_MISMATCH") from exc
    try:
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except Exception as exc:
        raise SourceIntakeFailure("FILE_STRUCTURE_REJECTED") from exc
    try:
        sheet_names = tuple(workbook.sheetnames)
        if not set(template.required_sheet_names) <= set(sheet_names) or not set(
            sheet_names
        ) <= set(template.allowed_sheet_names):
            raise SourceIntakeFailure("FILE_SHEET_REJECTED")
        all_rows: list[tuple[int, tuple[object, ...]]] = []
        row_offset = 0
        for sheet_name in sheet_names:
            worksheet = workbook[sheet_name]
            if worksheet.sheet_state != "visible":
                raise SourceIntakeFailure("FILE_HIDDEN_SHEET_REJECTED")
            iterator = worksheet.iter_rows()
            try:
                header_cells = next(iterator)
            except StopIteration as exc:
                raise SourceIntakeFailure("FILE_STRUCTURE_REJECTED") from exc
            _headers(template, (cell.value for cell in header_cells))
            for number, cells in enumerate(iterator, start=2):
                if any(cell.data_type == "f" for cell in cells):
                    raise SourceIntakeFailure("FILE_FORMULA_REJECTED")
                all_rows.append((row_offset + number, tuple(cell.value for cell in cells)))
            row_offset += max(worksheet.max_row, 1)
        return _typed_rows(template, all_rows)
    finally:
        workbook.close()


def validate_file(
    template: FileImportTemplateVersion, content: bytes, media_type: str
) -> ValidatedFileImport:
    if media_type not in template.accepted_media_types:
        raise SourceIntakeFailure("FILE_MEDIA_TYPE_REJECTED")
    if not content or len(content) > template.file_size_limit:
        raise SourceIntakeFailure("FILE_SIZE_REJECTED")
    if media_type == "text/csv":
        rows, rejected = _csv_rows(template, content)
    elif media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        rows, rejected = _xlsx_rows(template, content)
    else:
        raise SourceIntakeFailure("FILE_MEDIA_TYPE_REJECTED")
    return ValidatedFileImport(
        template_id=template.template_id,
        template_version=template.version,
        media_type=media_type,
        content_hash=hashlib.sha256(content).hexdigest(),
        schema_fingerprint=_schema_fingerprint(template),
        rows=rows,
        rejected_rows=rejected,
    )
